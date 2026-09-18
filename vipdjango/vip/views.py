import os
import re
import io
import csv
import logging
import uuid
from datetime import timedelta
from pathlib import Path
from urllib.parse import urlencode

from django.conf import settings
from django.http import HttpResponse, HttpResponseBadRequest, StreamingHttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.contrib.auth.decorators import login_required
from django.contrib.auth import get_user_model
from django.contrib.auth.forms import PasswordChangeForm
from django.contrib.auth.models import Group
from django.contrib.auth import update_session_auth_hash
from django.db import transaction
from django.db.models import Count, Max, Q
from django.urls import reverse
from django.utils.text import slugify
from django.utils import timezone
from django.views.decorators.http import require_POST

from .forms import (
    ClassGroupCreateForm,
    PromptTextUploadForm,
    RolePromptForm,
    StudentAccountCreateForm,
    StudentBulkUploadForm,
)
from .models import ChatMessage, ChatSession, RolePrompt
from .conversation_graph import MAX_TURNS
from .conversation_engine import (
    DEFAULT_INTRODUCTION,
    ConversationEngine,
    split_dialogue_and_voice,
)
from .conversation_scenario import parse_scenario_prompt

logger = logging.getLogger(__name__)

def _prompt_file_path(filename):
    candidates = [
        Path(settings.BASE_DIR).parent / "prompts" / filename,
        Path(settings.BASE_DIR) / "prompts" / filename,
    ]
    return next((path for path in candidates if path.exists()), None)


def _extract_template_prefix():
    global_prompt_path = _prompt_file_path("global_prompt.md")
    if global_prompt_path:
        return global_prompt_path.read_text(encoding="utf-8").strip()

    # Keep compatibility with legacy deployments that still provide the old
    # template file, but never prefer it over the global prompt.
    prompt_template_path = _prompt_file_path("prompt_template.md")
    if not prompt_template_path:
        return ""
    prompt_template = prompt_template_path.read_text(encoding="utf-8")
    if "{role}" in prompt_template:
        return prompt_template.split("{role}", 1)[0].strip()
    return prompt_template.strip()


def _load_api_key():
    api_key = os.getenv("OPENAI_API_KEY", "").strip()
    return api_key


def _clean_for_tts(text):
    return re.sub(r"\([^)]*\)", "", text).strip()


def _generate_assistant_response(
    role_text,
    session_messages,
    session=None,
):
    api_key = _load_api_key()
    if not api_key:
        raise RuntimeError("OPENAI_API_KEY is not configured")
    engine = ConversationEngine(_extract_template_prefix(), api_key)
    conversation_state = None
    if session is not None:
        conversation_state = {
            "core_question_state": session.core_question_state,
            "completion_status": session.completion_status,
            "current_stage": session.conversation_stage,
            "active_objective": session.active_objective,
            "covered_objectives": list(session.covered_objectives or []),
            "unresolved_objectives": list(session.unresolved_objectives or []),
            "recent_topics": list(session.recent_topics or []),
            "active_topic": session.active_topic,
            "covered_topics": list(session.covered_topics or []),
            "unresolved_topics": list(session.unresolved_topics or []),
            "topic_turn_counts": dict(session.topic_turn_counts or {}),
            "ending_ready": session.ending_ready,
        }
    return engine.respond(
        role_text,
        list(session_messages),
        conversation_state=conversation_state,
    )


def _request_turn_id(value):
    """Accept a safe client retry ID or create one for a non-voice submit."""
    candidate = (value or "").strip()
    if candidate and re.fullmatch(r"[A-Za-z0-9_-]{1,64}", candidate):
        return candidate
    return uuid.uuid4().hex


def _accept_learner_turn(session, user_text, turn_id):
    """Atomically claim one learner turn for a session.

    The session row is the lock. A retry with the same ID is idempotent; a
    different request cannot add a second pending learner message while the
    first response is being generated.
    """
    with transaction.atomic():
        locked_session = ChatSession.objects.select_for_update().get(pk=session.pk)
        existing_student = locked_session.messages.filter(
            sender=ChatMessage.Sender.STUDENT,
            turn_id=turn_id,
        ).first()
        if locked_session.ended_at is not None:
            return locked_session, "closed", None
        if locked_session.active_turn_id and (
            locked_session.active_claimed_at is None
            or locked_session.active_claimed_at < timezone.now() - timedelta(seconds=90)
        ):
            locked_session.active_turn_id = ""
            locked_session.active_claim_id = ""
            locked_session.active_claimed_at = None
            locked_session.save(update_fields=["active_turn_id", "active_claim_id", "active_claimed_at"])
        if existing_student:
            if existing_student.content.strip() != user_text.strip():
                logger.warning(
                    "Rejecting reused turn ID with different content session=%s turn=%s",
                    session.pk,
                    turn_id,
                )
                return locked_session, "conflict", None
            existing_assistant = locked_session.messages.filter(
                sender=ChatMessage.Sender.ASSISTANT,
                turn_id=turn_id,
            ).first()
            if existing_assistant or locked_session.last_completed_turn_id == turn_id:
                return locked_session, "duplicate", existing_assistant
            if locked_session.active_turn_id:
                return locked_session, "pending", None

            # A failed/cancelled request leaves its learner message in the
            # transcript so a retry preserves history. Reclaim that same
            # message instead of inserting a second learner turn.
            locked_session.active_turn_id = turn_id
            locked_session.active_claim_id = uuid.uuid4().hex
            locked_session.active_claimed_at = timezone.now()
            locked_session.save(update_fields=["active_turn_id", "active_claim_id", "active_claimed_at"])
            return locked_session, "accepted", None

        if locked_session.active_turn_id:
            return locked_session, "pending", None

        latest = locked_session.messages.order_by("-id").first()
        if latest and latest.sender == ChatMessage.Sender.STUDENT:
            return locked_session, "pending", None

        ChatMessage.objects.create(
            session=locked_session,
            sender=ChatMessage.Sender.STUDENT,
            content=user_text,
            turn_id=turn_id,
        )
        locked_session.active_turn_id = turn_id
        locked_session.active_claim_id = uuid.uuid4().hex
        locked_session.active_claimed_at = timezone.now()
        locked_session.save(update_fields=["active_turn_id", "active_claim_id", "active_claimed_at"])
        return locked_session, "accepted", None


def _release_learner_turn(session, turn_id, claim_id=""):
    """Release a failed/cancelled claim without deleting the learner row."""
    if not turn_id:
        return False
    with transaction.atomic():
        locked_session = ChatSession.objects.select_for_update().get(pk=session.pk)
        if locked_session.active_turn_id != turn_id:
            return False
        if claim_id and locked_session.active_claim_id != claim_id:
            return False
        locked_session.active_turn_id = ""
        locked_session.active_claim_id = ""
        locked_session.active_claimed_at = None
        locked_session.save(update_fields=["active_turn_id", "active_claim_id", "active_claimed_at"])
    return True


def _persist_assistant_response(
    session,
    assistant_text,
    conversation_complete,
    debug_info,
    turn_id="",
    cancellation_callback=None,
    claim_id="",
):
    """Commit at most one response, and never commit for a stale turn."""
    with transaction.atomic():
        locked_session = ChatSession.objects.select_for_update().get(pk=session.pk)
        if turn_id:
            if locked_session.ended_at is not None:
                return False
            if cancellation_callback and cancellation_callback():
                logger.info("Ignoring cancelled assistant response session=%s turn=%s", session.pk, turn_id)
                return False
            existing = locked_session.messages.filter(
                sender=ChatMessage.Sender.ASSISTANT,
                turn_id=turn_id,
            ).first()
            if existing:
                return False
            if locked_session.active_turn_id != turn_id:
                logger.warning("Ignoring unclaimed assistant response session=%s turn=%s", session.pk, turn_id)
                return False
            if claim_id and locked_session.active_claim_id != claim_id:
                logger.warning("Ignoring stale assistant claim session=%s turn=%s", session.pk, turn_id)
                return False
            latest = locked_session.messages.order_by("-id").first()
            if not latest or latest.sender != ChatMessage.Sender.STUDENT or latest.turn_id != turn_id:
                logger.warning("Ignoring stale assistant response session=%s turn=%s", session.pk, turn_id)
                return False

        if assistant_text:
            ChatMessage.objects.create(
                session=locked_session,
                sender=ChatMessage.Sender.ASSISTANT,
                content=assistant_text,
                voice_metadata=debug_info.get("voice_metadata", ""),
                turn_id=turn_id,
            )
        locked_session.conversation_stage = debug_info.get("current_stage", locked_session.conversation_stage)
        locked_session.conversation_phase = debug_info.get("phase", locked_session.conversation_phase)
        locked_session.completion_status = conversation_complete
        locked_session.active_objective = debug_info.get("active_objective", locked_session.active_objective)
        covered_objectives = debug_info.get("covered_objectives")
        if isinstance(covered_objectives, list):
            locked_session.covered_objectives = covered_objectives
        unresolved_objectives = debug_info.get("unresolved_objectives")
        if isinstance(unresolved_objectives, list):
            locked_session.unresolved_objectives = unresolved_objectives
        recent_topics = debug_info.get("recent_topics")
        if isinstance(recent_topics, list):
            locked_session.recent_topics = recent_topics
        locked_session.active_topic = debug_info.get("active_topic", locked_session.active_topic)
        covered_topics = debug_info.get("covered_topics")
        if isinstance(covered_topics, list):
            locked_session.covered_topics = covered_topics
        unresolved_topics = debug_info.get("unresolved_topics")
        if isinstance(unresolved_topics, list):
            locked_session.unresolved_topics = unresolved_topics
        topic_turn_counts = debug_info.get("topic_turn_counts")
        if isinstance(topic_turn_counts, dict):
            locked_session.topic_turn_counts = topic_turn_counts
        if "ending_ready" in debug_info:
            locked_session.ending_ready = bool(debug_info["ending_ready"])
        update_fields = [
            "conversation_stage",
            "conversation_phase",
            "completion_status",
            "active_objective",
            "covered_objectives",
            "unresolved_objectives",
            "recent_topics",
            "active_topic",
            "covered_topics",
            "unresolved_topics",
            "topic_turn_counts",
            "ending_ready",
        ]
        if turn_id:
            locked_session.active_turn_id = ""
            locked_session.active_claim_id = ""
            locked_session.active_claimed_at = None
            locked_session.last_completed_turn_id = turn_id
            update_fields.extend(["active_turn_id", "active_claim_id", "active_claimed_at", "last_completed_turn_id"])
        if "core_question_state" in debug_info:
            locked_session.core_question_state = debug_info["core_question_state"]
            update_fields.append("core_question_state")
        if conversation_complete and locked_session.ended_at is None:
            locked_session.ended_at = timezone.now()
            update_fields.append("ended_at")
        locked_session.save(update_fields=update_fields)
    return True


def _clean_name_part(value):
    return re.sub(r"[^a-zA-Z0-9]", "", value or "").lower()


def _normalize_net_id(value):
    return (value or "").strip().lower()


def _build_student_password(first_name, last_name, student_id):
    return f"{_clean_name_part(first_name)}{_clean_name_part(last_name)}{student_id}"


def _normalize_class_name(value):
    return re.sub(r"\s+", " ", (value or "").strip())


def _normalize_header(header):
    return re.sub(r"[^a-z0-9]", "", (header or "").strip().lower())


def _extract_row_value(row, aliases):
    normalized = {_normalize_header(k): v for k, v in row.items()}
    for alias in aliases:
        value = normalized.get(alias)
        if value is not None and str(value).strip() != "":
            return str(value).strip()
    return ""


def _read_bulk_rows(uploaded_file):
    suffix = Path(uploaded_file.name).suffix.lower()
    if suffix == ".csv":
        decoded = uploaded_file.read().decode("utf-8-sig")
        return list(csv.DictReader(io.StringIO(decoded)))

    if suffix == ".xlsx":
        try:
            import openpyxl
        except ImportError as exc:
            raise ValueError("openpyxl is required to read .xlsx files.") from exc

        wb = openpyxl.load_workbook(uploaded_file, data_only=True)
        ws = wb.active
        headers = [str(cell.value).strip() if cell.value is not None else "" for cell in ws[1]]
        rows = []
        for values in ws.iter_rows(min_row=2, values_only=True):
            if all(v is None or str(v).strip() == "" for v in values):
                continue
            rows.append(
                {
                    headers[i]: ("" if values[i] is None else str(values[i]).strip())
                    for i in range(len(headers))
                }
            )
        return rows

    raise ValueError("Unsupported file type. Please upload .xlsx or .csv.")


def _create_student_account(first_name, last_name, net_id, student_id, class_name):
    User = get_user_model()
    username = _normalize_net_id(net_id)
    password = _build_student_password(first_name, last_name, student_id)
    normalized_class = _normalize_class_name(class_name)

    if not username:
        return None, "NetID is required."
    if User.objects.filter(username=username).exists():
        return None, f'NetID "{username}" already exists.'

    user = User.objects.create_user(
        username=username,
        password=password,
        first_name=first_name.strip(),
        last_name=last_name.strip(),
    )

    student_group, _ = Group.objects.get_or_create(name="Student")
    user.groups.add(student_group)
    if normalized_class:
        class_group, _ = Group.objects.get_or_create(name=f"Class: {normalized_class}")
        user.groups.add(class_group)
    return {
        "username": username,
        "password": password,
        "net_id": username,
        "first_name": first_name.strip(),
        "last_name": last_name.strip(),
        "student_id": student_id.strip(),
        "class_name": normalized_class or "Unassigned",
    }, None


def _is_professor(user):
    return user.is_superuser or user.groups.filter(name__iexact="Professor").exists()


def _is_student(user):
    if _is_professor(user):
        return False
    return (
        user.groups.filter(name__startswith="Class: ")
        .exclude(name__iexact="Class: Unassigned")
        .exists()
    )


def _student_queryset():
    User = get_user_model()
    return (
        User.objects.filter(Q(groups__name__startswith="Class: ") | Q(groups__name="Student"))
        .exclude(groups__name="Professor")
        .distinct()
    )


@login_required
def home(request):
    if _is_professor(request.user):
        return redirect("vip:professor_dashboard")
    elif _is_student(request.user):
        return redirect("vip:student_dashboard")
    else:
        return redirect("vip:dashboard")
    
@login_required
def dashboard(request):
    return render(request, "vip/dashboard.html")


@login_required
def account_settings(request):
    is_professor = _is_professor(request.user)
    status_message = ""
    error_message = ""
    password_form = PasswordChangeForm(request.user)

    if request.method == "POST":
        action = request.POST.get("action")

        if action == "change_password":
            password_form = PasswordChangeForm(request.user, request.POST)
            if password_form.is_valid():
                user = password_form.save()
                update_session_auth_hash(request, user)
                status_message = "Password updated successfully."
            else:
                error_message = "Please fix password form errors."

        if action == "update_name" and is_professor:
            first_name = request.POST.get("first_name", "").strip()
            last_name = request.POST.get("last_name", "").strip()
            request.user.first_name = first_name
            request.user.last_name = last_name
            request.user.save(update_fields=["first_name", "last_name"])
            status_message = "Profile name updated."

    return render(
        request,
        "vip/account_settings.html",
        {
            "is_professor": is_professor,
            "status_message": status_message,
            "error_message": error_message,
            "password_form": password_form,
        },
    )

# Professor views
@login_required
def professor_dashboard(request):
    if not _is_professor(request.user):
        return redirect("vip:home")

    current_tab = request.GET.get("tab", "prompts")
    if current_tab not in {"prompts", "logs", "test_chat"}:
        current_tab = "prompts"

    prompts = RolePrompt.objects.order_by("-updated_at")
    prompt_upload_form = PromptTextUploadForm()
    students = (
        _student_queryset().prefetch_related("groups")
        .annotate(
            session_count=Count("chat_sessions", distinct=True),
            last_session_at=Max("chat_sessions__started_at"),
        )
        .order_by("username")
    )
    students_by_class = {}
    for student in students:
        class_names = []
        for group in student.groups.all():
            if group.name.startswith("Class: ") and group.name.lower() != "class: unassigned":
                class_names.append(group.name.replace("Class: ", "", 1))
        if not class_names:
            class_names = ["Unassigned"]
        for class_name in class_names:
            students_by_class.setdefault(class_name, []).append(student)
    students_by_class = [
        {"class_name": class_name, "students": students_by_class[class_name]}
        for class_name in sorted(students_by_class.keys())
    ]
    return render(
        request,
        "vip/professor_dashboard.html",
        {
            "prompts": prompts,
            "prompt_upload_form": prompt_upload_form,
            "students_by_class": students_by_class,
            "current_tab": current_tab,
        },
    )


@login_required
def professor_download_prompt_txt(request, prompt_id):
    if not _is_professor(request.user):
        return redirect("vip:home")

    prompt = get_object_or_404(RolePrompt, pk=prompt_id)
    filename_base = slugify(prompt.title) or f"prompt-{prompt.id}"
    response = HttpResponse(prompt.content, content_type="text/plain; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="{filename_base}.txt"'
    return response


@login_required
def professor_download_prompt_template(request):
    if not _is_professor(request.user):
        return redirect("vip:home")

    template_path = _prompt_file_path("role_prompt.md")
    if not template_path:
        return HttpResponseBadRequest("Template file not found.")

    content = template_path.read_text(encoding="utf-8")
    response = HttpResponse(content, content_type="text/plain; charset=utf-8")
    response["Content-Disposition"] = 'attachment; filename="role_prompt_template.md"'
    return response


@login_required
@require_POST
def professor_upload_prompt_file(request):
    if not _is_professor(request.user):
        return redirect("vip:home")

    form = PromptTextUploadForm(request.POST, request.FILES)
    if not form.is_valid():
        return redirect(f"{reverse('vip:professor_dashboard')}?tab=prompts")

    uploaded = form.cleaned_data["file"]
    try:
        content = uploaded.read().decode("utf-8-sig")
    except UnicodeDecodeError:
        return HttpResponseBadRequest("File must be UTF-8 text.")

    title = form.cleaned_data["title"].strip() if form.cleaned_data["title"] else Path(uploaded.name).stem
    parsed_initial = RolePromptForm.initial_from_content(
        content,
        title=title or "Uploaded Prompt",
        is_active=form.cleaned_data["is_active"],
    )
    structured_form = RolePromptForm(parsed_initial)
    if structured_form.is_valid():
        structured_content = structured_form.render_markdown_content()
    else:
        structured_content = content

    RolePrompt.objects.create(
        title=parsed_initial["title"] or "Uploaded Prompt",
        content=structured_content,
        created_by=request.user,
        is_active=form.cleaned_data["is_active"],
    )
    return redirect(f"{reverse('vip:professor_dashboard')}?tab=prompts")


@login_required
def professor_manage_accounts(request):
    if not _is_professor(request.user):
        return redirect("vip:home")

    legacy_unassigned_group = Group.objects.filter(name__iexact="Class: Unassigned").first()
    if legacy_unassigned_group:
        legacy_unassigned_group.user_set.clear()
        legacy_unassigned_group.delete()

    single_form = StudentAccountCreateForm()
    bulk_form = StudentBulkUploadForm()
    class_form = ClassGroupCreateForm()
    created_accounts = []
    skipped_rows = []
    info_message = ""

    if request.method == "POST":
        action = request.POST.get("action")

        if action == "create_single":
            single_form = StudentAccountCreateForm(request.POST)
            if single_form.is_valid():
                selected_class = _normalize_class_name(single_form.cleaned_data["class_name"])
                class_group = None
                if selected_class:
                    class_group = Group.objects.filter(name=f"Class: {selected_class}").first()
                if selected_class and not class_group:
                    skipped_rows.append(
                        {"row": "Single form", "reason": "Please select an existing class group from the dropdown."}
                    )
                else:
                    account, error = _create_student_account(
                        first_name=single_form.cleaned_data["first_name"],
                        last_name=single_form.cleaned_data["last_name"],
                        net_id=single_form.cleaned_data["net_id"],
                        student_id=single_form.cleaned_data["student_id"],
                        class_name=selected_class,
                    )
                    if error:
                        skipped_rows.append({"row": "Single form", "reason": error})
                    else:
                        created_accounts.append(account)
                        info_message = "Student account created."

        if action == "bulk_upload":
            bulk_form = StudentBulkUploadForm(request.POST, request.FILES)
            if bulk_form.is_valid():
                aliases = {
                    "first_name": ["firstname", "first", "givenname"],
                    "last_name": ["lastname", "last", "surname", "familyname"],
                    "net_id": ["netid", "net_id", "username", "login"],
                    "student_id": ["studentid", "id", "studentnumber", "sid"],
                    "class_name": ["classname", "class", "course", "section", "group"],
                }
                try:
                    rows = _read_bulk_rows(bulk_form.cleaned_data["file"])
                except ValueError as exc:
                    skipped_rows.append({"row": "Upload", "reason": str(exc)})
                    rows = []

                for idx, row in enumerate(rows, start=2):
                    first_name = _extract_row_value(row, aliases["first_name"])
                    last_name = _extract_row_value(row, aliases["last_name"])
                    net_id = _extract_row_value(row, aliases["net_id"])
                    student_id = _extract_row_value(row, aliases["student_id"])
                    class_name = _extract_row_value(row, aliases["class_name"])
                    if re.match(r"^\d+\.0$", student_id):
                        student_id = student_id[:-2]

                    if not first_name or not last_name or not net_id or not student_id:
                        skipped_rows.append(
                            {
                                "row": idx,
                                "reason": "Missing required fields (first_name, last_name, net_id, student_id).",
                            }
                        )
                        continue
                    if not student_id.isdigit():
                        skipped_rows.append({"row": idx, "reason": "student_id must be numeric."})
                        continue

                    account, error = _create_student_account(
                        first_name=first_name,
                        last_name=last_name,
                        net_id=net_id,
                        student_id=student_id,
                        class_name=class_name,
                    )
                    if error:
                        skipped_rows.append({"row": idx, "reason": error})
                    else:
                        created_accounts.append(account)

                info_message = f"Bulk upload complete. Created {len(created_accounts)} account(s)."

        if action == "create_class":
            class_form = ClassGroupCreateForm(request.POST)
            if class_form.is_valid():
                normalized_class = _normalize_class_name(class_form.cleaned_data["class_name"])
                if not normalized_class:
                    skipped_rows.append({"row": "Create class", "reason": "Class name cannot be blank."})
                else:
                    Group.objects.get_or_create(name=f"Class: {normalized_class}")
                    info_message = f'Class group "{normalized_class}" is ready.'

        if action == "move_student":
            student_id = request.POST.get("student_id", "").strip()
            selected_class_group_id = request.POST.get("class_group_id", "").strip()
            student = get_object_or_404(_student_queryset(), pk=student_id)
            existing_class_groups = student.groups.filter(name__startswith="Class: ")
            student.groups.remove(*existing_class_groups)

            if selected_class_group_id == "__UNASSIGNED__":
                info_message = f"Moved {student.username} to Unassigned."
            elif selected_class_group_id:
                class_group = Group.objects.filter(pk=selected_class_group_id, name__startswith="Class: ").first()
                if class_group:
                    student.groups.add(class_group)
                    class_name = class_group.name.replace("Class: ", "", 1)
                    info_message = f"Moved {student.username} to class {class_name}."
                else:
                    skipped_rows.append({"row": "Move student", "reason": "Selected class group does not exist."})
            else:
                skipped_rows.append({"row": "Move student", "reason": "Please choose a class from the dropdown."})

        if action == "delete_student":
            student_id = request.POST.get("student_id", "").strip()
            student = get_object_or_404(_student_queryset(), pk=student_id)
            username = student.username
            student.delete()
            info_message = f"Deleted student account {username}."

        if action == "delete_class":
            class_group_id = request.POST.get("class_group_id", "").strip()
            class_group = Group.objects.filter(pk=class_group_id, name__startswith="Class: ").first()
            if class_group:
                class_name = class_group.name.replace("Class: ", "", 1)
                class_group.delete()
                info_message = f'Deleted class group "{class_name}".'
            else:
                skipped_rows.append({"row": "Delete class", "reason": "Class group not found."})

    class_groups = (
        Group.objects.filter(name__startswith="Class: ")
        .exclude(name__iexact="Class: Unassigned")
        .order_by("name")
    )
    roster_students = _student_queryset().prefetch_related("groups").order_by("username")
    student_rows = []
    for student in roster_students:
        class_names = [
            group.name.replace("Class: ", "", 1)
            for group in student.groups.all()
            if group.name.startswith("Class: ") and group.name.lower() != "class: unassigned"
        ]
        student_rows.append(
            {
                "id": student.id,
                "username": student.username,
                "first_name": student.first_name,
                "last_name": student.last_name,
                "classes_display": ", ".join(class_names) if class_names else "Unassigned",
            }
        )

    return render(
        request,
        "vip/professor_accounts.html",
        {
            "single_form": single_form,
            "bulk_form": bulk_form,
            "class_form": class_form,
            "created_accounts": created_accounts,
            "skipped_rows": skipped_rows,
            "info_message": info_message,
            "class_groups": class_groups,
            "student_rows": student_rows,
        },
    )


@login_required
def create_prompt(request):
    if not _is_professor(request.user):
        return redirect("vip:home")

    if request.method == "POST":
        form = RolePromptForm(request.POST)
        if form.is_valid():
            RolePrompt.objects.create(
                title=form.cleaned_data["title"],
                content=form.render_markdown_content(),
                created_by=request.user,
                is_active=form.cleaned_data["is_active"],
            )
            return redirect("vip:professor_dashboard")
    else:
        form = RolePromptForm()

    return render(
        request,
        "vip/prompt_form.html",
        {"form": form, "page_title": "Add Prompt"},
    )

@login_required
def edit_prompt(request, prompt_id):
    if not _is_professor(request.user):
        return redirect("vip:home")

    prompt = get_object_or_404(RolePrompt, pk=prompt_id)

    if request.method == "POST":
        form = RolePromptForm(request.POST)
        if form.is_valid():
            prompt.title = form.cleaned_data["title"]
            prompt.is_active = form.cleaned_data["is_active"]
            prompt.content = form.render_markdown_content()
            prompt.save(update_fields=["title", "is_active", "content", "updated_at"])
            return redirect("vip:professor_dashboard")
    else:
        form = RolePromptForm(initial=RolePromptForm.initial_from_prompt(prompt))

    return render(
        request,
        "vip/prompt_form.html",
        {"form": form, "page_title": "Edit Prompt"},
    )


@login_required
@require_POST
def set_active_prompt(request, prompt_id):
    if not _is_professor(request.user):
        return redirect("vip:home")

    selected_prompt = get_object_or_404(RolePrompt, pk=prompt_id)
    selected_prompt.is_active = True
    selected_prompt.save(update_fields=["is_active", "updated_at"])
    return redirect("vip:professor_dashboard")


@login_required
@require_POST
def deactivate_prompt(request, prompt_id):
    if not _is_professor(request.user):
        return redirect("vip:home")

    selected_prompt = get_object_or_404(RolePrompt, pk=prompt_id)
    selected_prompt.is_active = False
    selected_prompt.save(update_fields=["is_active", "updated_at"])
    return redirect("vip:professor_dashboard")


@login_required
@require_POST
def delete_prompt(request, prompt_id):
    if not _is_professor(request.user):
        return redirect("vip:home")

    prompt = get_object_or_404(RolePrompt, pk=prompt_id)
    # Close any active sessions using this prompt before deletion so students
    # don't continue a conversation with a removed prompt.
    ChatSession.objects.filter(role_prompt=prompt, ended_at__isnull=True).update(ended_at=timezone.now())
    prompt.delete()
    return redirect(f"{reverse('vip:professor_dashboard')}?tab=prompts")


@login_required
def professor_session_detail(request, session_id):
    if not _is_professor(request.user):
        return redirect("vip:home")

    session = get_object_or_404(
        ChatSession.objects.select_related("student", "role_prompt"),
        pk=session_id,
    )
    messages = session.messages.order_by("created_at")

    return render(
        request,
        "vip/professor_session_detail.html",
        {
            "session": session,
            "messages": messages,
            "rendered_messages": _rendered_chat_messages(session, "Student"),
        },
    )


@login_required
def professor_student_logs(request, student_id):
    if not _is_professor(request.user):
        return redirect("vip:home")

    student = get_object_or_404(_student_queryset(), pk=student_id)
    sessions = (
        ChatSession.objects.filter(student=student)
        .select_related("role_prompt")
        .prefetch_related("messages")
        .order_by("-started_at")
    )

    return render(
        request,
        "vip/professor_student_logs.html",
        {
            "student": student,
            "sessions": sessions,
        },
    )


@login_required
@require_POST
def professor_delete_session(request, session_id):
    if not _is_professor(request.user):
        return redirect("vip:home")

    session = get_object_or_404(ChatSession, pk=session_id)
    student_id = session.student_id
    session.delete()

    next_url = request.POST.get("next")
    if next_url:
        return redirect(next_url)
    return redirect("vip:professor_student_logs", student_id=student_id)


@login_required
@require_POST
def professor_delete_student_logs(request, student_id):
    if not _is_professor(request.user):
        return redirect("vip:home")

    student = get_object_or_404(_student_queryset(), pk=student_id)
    ChatSession.objects.filter(student=student).delete()
    return redirect("vip:professor_student_logs", student_id=student_id)


@login_required
@require_POST
def professor_reset_all_student_logs(request):
    if not _is_professor(request.user):
        return redirect("vip:home")

    student_ids = _student_queryset().values_list("id", flat=True)
    ChatSession.objects.filter(student_id__in=student_ids).delete()
    return redirect(f"{reverse('vip:professor_dashboard')}?tab=logs")


def _chat_url(view_name, selected_prompt=None, **params):
    query = {}
    if selected_prompt:
        query["prompt"] = selected_prompt.id
    query.update({key: value for key, value in params.items() if value is not None})
    url = reverse(view_name)
    return f"{url}?{urlencode(query)}" if query else url


def _rendered_chat_messages(session, user_label):
    if not session:
        return []

    rendered = []
    for message in session.messages.order_by("created_at"):
        display_content = message.content
        embedded_emotion = ""
        if message.sender == ChatMessage.Sender.ASSISTANT:
            display_content, embedded_emotion = split_dialogue_and_voice(message.content)
        rendered.append(
            {
                "id": message.id,
                "sender": message.sender,
                "sender_display": "AI" if message.sender == ChatMessage.Sender.ASSISTANT else user_label,
                "created_at": message.created_at,
                "display_content": display_content,
                "voice_enabled": bool(message.voice_metadata or embedded_emotion),
            }
        )
    return rendered


def _chat_dashboard_context(
    active_prompts,
    selected_prompt,
    sessions,
    current_session,
    error_message,
    user_label,
    force_new,
):
    latest = current_session.messages.order_by("-id").first() if current_session else None
    pending = latest if latest and latest.sender == ChatMessage.Sender.STUDENT else None
    return {
        "active_prompts": active_prompts,
        "selected_prompt": selected_prompt,
        "sessions": sessions,
        "current_session": current_session,
        "rendered_messages": _rendered_chat_messages(current_session, user_label),
        "error_message": error_message,
        "force_new": force_new,
        "pending_message": pending,
    }


def _selected_chat_state(request):
    active_prompts = RolePrompt.objects.filter(is_active=True).order_by("title")
    sessions = (
        ChatSession.objects.filter(student=request.user)
        .select_related("role_prompt")
        .prefetch_related("messages")
        .order_by("-started_at")
    )

    selected_prompt = None
    selected_prompt_id = request.POST.get("prompt_id") or request.GET.get("prompt")
    if selected_prompt_id:
        selected_prompt = active_prompts.filter(pk=selected_prompt_id).first()
    if not selected_prompt:
        selected_prompt = active_prompts.first()

    current_session = None
    session_id = request.POST.get("session_id") or request.GET.get("session")
    if session_id:
        current_session = get_object_or_404(
            ChatSession.objects.filter(student=request.user).select_related("role_prompt"),
            pk=session_id,
        )
        if selected_prompt_id and current_session.role_prompt_id != getattr(selected_prompt, "id", None):
            current_session = sessions.filter(role_prompt=selected_prompt).first()
        elif current_session.role_prompt and not selected_prompt_id:
            selected_prompt = current_session.role_prompt
    elif request.GET.get("new") != "1" and selected_prompt:
        current_session = sessions.filter(role_prompt=selected_prompt).first()

    force_new = request.POST.get("force_new") == "1" or request.GET.get("force_new") == "1"
    return active_prompts, sessions, selected_prompt, current_session, force_new


def _usable_chat_session(user, selected_prompt, current_session, force_new, turn_id=""):
    needs_session = (
        not current_session
        or current_session.role_prompt_id != selected_prompt.id
        or (current_session.ended_at is not None and not turn_id)
        or force_new
    )
    if needs_session:
        current_session = None
        if not force_new:
            current_session = (
                ChatSession.objects.filter(
                    student=user,
                    role_prompt=selected_prompt,
                    ended_at__isnull=True,
                )
                .order_by("-started_at")
                .first()
            )
        return current_session or _create_chat_session(user, selected_prompt), None

    return current_session, None


def _seed_introduction(session):
    """Persist the fixed Introduction as text when a new chat starts."""
    if not session or not session.role_prompt:
        return
    if session.messages.filter(sender=ChatMessage.Sender.ASSISTANT).exists():
        return

    scenario = parse_scenario_prompt(session.scenario_content or session.role_prompt.content)
    introduction = (scenario.introduction or DEFAULT_INTRODUCTION).strip()
    if introduction:
        ChatMessage.objects.create(
            session=session,
            sender=ChatMessage.Sender.ASSISTANT,
            content=introduction,
        )


def _create_chat_session(user, selected_prompt):
    session = ChatSession.objects.create(student=user, role_prompt=selected_prompt, scenario_content=selected_prompt.content)
    _seed_introduction(session)
    return session


def _chat_dashboard(
    request,
    *,
    template_name,
    view_name,
    user_label,
    no_prompt_message,
    allow_delete_session=False,
):
    active_prompts, sessions, selected_prompt, current_session, force_new = _selected_chat_state(request)
    error_message = None

    if request.method == "POST":
        action = request.POST.get("action")
        if action == "new_session":
            if selected_prompt:
                session = _create_chat_session(request.user, selected_prompt)
                return redirect(_chat_url(view_name, selected_prompt, session=session.id))
            return redirect(_chat_url(view_name, selected_prompt))

        if allow_delete_session and action == "delete_session":
            target_session_id = request.POST.get("target_session_id") or request.POST.get("session_id")
            if target_session_id:
                get_object_or_404(ChatSession, pk=target_session_id, student=request.user).delete()
            return redirect(_chat_url(view_name, selected_prompt, new=1))

        if action == "close_conversation":
            if current_session and current_session.ended_at is None:
                _release_learner_turn(
                    current_session,
                    current_session.active_turn_id,
                    current_session.active_claim_id,
                )
                current_session.ended_at = timezone.now()
                current_session.save(update_fields=["ended_at"])
            return redirect(_chat_url(view_name, selected_prompt, session=current_session.id) if current_session else _chat_url(view_name, selected_prompt))

        if action == "send_message":
            user_text = request.POST.get("message", "").strip()
            requested_turn_id = request.POST.get("turn_id", "").strip()
            turn_id = _request_turn_id(requested_turn_id)
            retry_turn_id = turn_id if requested_turn_id == turn_id else ""
            if not selected_prompt:
                error_message = no_prompt_message
            elif not user_text:
                error_message = "Please type a message before sending."
            else:
                current_session, error_message = _usable_chat_session(
                    request.user,
                    selected_prompt,
                    current_session,
                    force_new,
                    turn_id=retry_turn_id,
                )

                if not error_message:
                    completed_turns = current_session.messages.filter(
                        sender=ChatMessage.Sender.STUDENT
                    ).count()
                    retry_exists = bool(
                        retry_turn_id
                        and current_session.messages.filter(turn_id=retry_turn_id).exists()
                    )
                    if completed_turns >= MAX_TURNS and not retry_exists:
                        error_message = "This conversation reached its safety limit before a natural ending."

            if error_message:
                return render(
                    request,
                    template_name,
                    _chat_dashboard_context(
                        active_prompts,
                        selected_prompt,
                        sessions,
                        current_session,
                        error_message,
                        user_label,
                        force_new,
                    ),
                )

            _seed_introduction(current_session)
            current_session, turn_status, existing_assistant = _accept_learner_turn(
                current_session,
                user_text,
                turn_id,
            )
            if turn_status == "pending":
                error_message = "Please wait for the AI response before sending another message."
            elif turn_status == "closed":
                return redirect(_chat_url(view_name, selected_prompt, session=current_session.id))
            elif turn_status == "conflict":
                error_message = "This turn ID was already used for different message content. Please retry from the chat page."

            if error_message:
                return render(
                    request,
                    template_name,
                    _chat_dashboard_context(
                        active_prompts,
                        selected_prompt,
                        sessions,
                        current_session,
                        error_message,
                        user_label,
                        force_new,
                    ),
                )

            if turn_status == "duplicate":
                return redirect(_chat_url(view_name, selected_prompt, session=current_session.id))

            try:
                assistant_text, conversation_complete, debug_info = _generate_assistant_response(
                    current_session.scenario_content or selected_prompt.content,
                    current_session.messages.order_by("created_at"),
                    current_session,
                )
            except Exception:
                _release_learner_turn(current_session, turn_id, current_session.active_claim_id)
                logger.exception("Chat response generation failed session=%s turn=%s", current_session.id, turn_id)
                return render(request, template_name, _chat_dashboard_context(
                    active_prompts, selected_prompt, sessions, current_session,
                    "The response could not be generated. Your message is saved. Press Retry response; if this continues, contact your instructor.",
                    user_label, False,
                ), status=503)
            _persist_assistant_response(
                current_session,
                assistant_text,
                conversation_complete,
                debug_info,
                turn_id=turn_id,
                claim_id=current_session.active_claim_id,
            )
            logger.debug(
                "Chat response generated: user=%s session=%s phase=%s complete=%s reason=%s ending_check=%s",
                request.user.username,
                current_session.id,
                debug_info.get("phase"),
                conversation_complete,
                debug_info.get("reason"),
                debug_info.get("ending_check"),
            )
            return redirect(
                _chat_url(
                    view_name,
                    selected_prompt,
                    session=current_session.id,
                    autoplay=1,
                )
            )

    return render(
        request,
        template_name,
        _chat_dashboard_context(
            active_prompts,
            selected_prompt,
            sessions,
            current_session,
            error_message,
            user_label,
            force_new,
        ),
    )


@login_required
def professor_test_chat(request):
    if not _is_professor(request.user):
        return redirect("vip:home")

    return _chat_dashboard(
        request,
        template_name="vip/professor_test_chat.html",
        view_name="vip:professor_test_chat",
        user_label="Professor",
        no_prompt_message="No scenario is available. Create or upload a scenario first.",
        allow_delete_session=True,
    )


# Student views

@login_required
def student_dashboard(request):
    if not _is_student(request.user):
        return redirect("vip:home")

    return _chat_dashboard(
        request,
        template_name="vip/student_dashboard.html",
        view_name="vip:student_dashboard",
        user_label="Student",
        no_prompt_message="No active prompt is available. Ask your professor to activate one.",
    )


@login_required
def student_download_session(request, session_id):
    if not (_is_student(request.user) or _is_professor(request.user)):
        return redirect("vip:home")

    session = get_object_or_404(
        ChatSession.objects.select_related("student", "role_prompt"),
        pk=session_id,
        student=request.user,
    )
    messages = session.messages.order_by("created_at")

    lines = [
        f"Student: {session.student.username}",
        f"Prompt: {session.role_prompt.title if session.role_prompt else 'None'}",
        f"Session ID: {session.id}",
        f"Started: {session.started_at}",
        "",
    ]
    for message in messages:
        speaker = "Student" if message.sender == ChatMessage.Sender.STUDENT else "Assistant"
        content = message.content
        voice_metadata = ""
        if message.sender == ChatMessage.Sender.ASSISTANT:
            content, embedded_voice = split_dialogue_and_voice(content)
            voice_metadata = (message.voice_metadata or embedded_voice).strip()
        lines.append(f"[{message.created_at}] {speaker}: {content}")
        if voice_metadata:
            lines.append(f"Voice metadata: {voice_metadata}")
        lines.append("")

    response = HttpResponse("\n".join(lines), content_type="text/plain")
    response["Content-Disposition"] = f'attachment; filename="chat_session_{session.id}.txt"'
    return response


@login_required
def student_message_tts(request, message_id):
    if not (_is_student(request.user) or _is_professor(request.user)):
        return redirect("vip:home")

    message = get_object_or_404(
        ChatMessage.objects.select_related("session"),
        pk=message_id,
        session__student=request.user,
    )
    if message.sender != ChatMessage.Sender.ASSISTANT:
        return HttpResponseBadRequest("TTS is only available for assistant messages.")

    dialogue, embedded_emotion = split_dialogue_and_voice(message.content)
    emotion = getattr(message, "voice_metadata", "") or embedded_emotion
    use_emotion_voice = request.GET.get("emotion", "1") != "0"
    if not emotion:
        return HttpResponseBadRequest("Audio is not available for this text-only assistant message.")

    api_key = _load_api_key()
    if not api_key:
        return HttpResponseBadRequest("API key is not configured.")

    try:
        from .speech import SpeechStream

        stream = SpeechStream(
            api_key,
            model=os.getenv("OPENAI_TTS_MODEL", "gpt-4o-mini-tts"),
            voice="onyx" if re.search(r"\bmale voice\b", emotion.lower()) else "coral",
            input=dialogue or "No spoken dialogue.",
            instructions=emotion if use_emotion_voice else "Speak naturally and clearly.",
            response_format="mp3",
        )
        result = StreamingHttpResponse(stream, content_type="audio/mpeg")
        result["Cache-Control"] = "no-store"
        result["X-Accel-Buffering"] = "no"
        return result
    except Exception:
        logger.exception("Speech generation failed message=%s", message.id)
        return HttpResponse("Voice is temporarily unavailable. You can continue using text.", status=503)
